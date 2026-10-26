import os
import csv
from typing import List, Dict, Any

from django.core.management.base import BaseCommand, CommandError

try:
    import openpyxl
except ImportError:
    openpyxl = None


class Command(BaseCommand):
    help = (
        "Extracts verified gold rows (company_id, category, correct_url, notes, verified=TRUE) "
        "from a checked company_review_single.xlsx workbook."
    )

    def add_arguments(self, parser):
        parser.add_argument(
            "--file",
            type=str,
            required=True,
            help="Path to the checked workbook (e.g. exports/crawl_review/20261023_112755/company_review_single.xlsx)",
        )
        parser.add_argument(
            "--out",
            type=str,
            default="tests/golden/gold_sources.csv",
            help="Destination output CSV path (default: tests/golden/gold_sources.csv)",
        )
        parser.add_argument(
            "--force",
            action="store_true",
            help="Force overwrite existing destination file without prompting.",
        )

    def handle(self, *args, **options):
        if openpyxl is None:
            raise CommandError("openpyxl is required to parse Excel workbooks. Please install openpyxl.")

        file_path = options["file"]
        if not os.path.isabs(file_path):
            file_path = os.path.abspath(file_path)

        if not os.path.isfile(file_path):
            raise CommandError(f"Workbook file does not exist: {file_path}")

        out_path = options["out"]
        if not os.path.isabs(out_path):
            out_path = os.path.abspath(out_path)

        # Rule: Never overwrite an existing file without asking (or explicit --force)
        if os.path.exists(out_path) and not options.get("force"):
            self.stdout.write(self.style.WARNING(f"File already exists at: {out_path}"))
            confirm = input("Destination file exists. Overwrite? (y/N): ").strip().lower()
            if confirm != "y":
                self.stdout.write(self.style.NOTICE("Aborted. Target file was not modified."))
                return

        wb = openpyxl.load_workbook(file_path, data_only=True)
        if "Review" not in wb.sheetnames:
            raise CommandError("Workbook is missing required 'Review' sheet.")

        ws = wb["Review"]
        headers = [cell.value for cell in ws[1]]

        def get_col_idx(name: str) -> int:
            try:
                return headers.index(name)
            except ValueError:
                raise CommandError(f"Missing column '{name}' in sheet 'Review'.")

        rec_type_idx = get_col_idx("record_type")
        comp_id_idx = get_col_idx("company_id")
        category_idx = get_col_idx("category")
        url_idx = get_col_idx("url")
        check_idx = get_col_idx("CHECK")
        correct_val_idx = get_col_idx("CORRECT_VALUE")
        notes_idx = get_col_idx("NOTES")

        gold_rows: List[Dict[str, Any]] = []

        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or not any(row):
                continue

            rec_type = str(row[rec_type_idx] or "").strip().upper()
            if rec_type != "PAGE":
                continue

            comp_id = str(row[comp_id_idx] or "").strip()
            category = str(row[category_idx] or "").strip()
            orig_url = str(row[url_idx] or "").strip()
            check_val = str(row[check_idx] or "").strip().capitalize()
            correct_val = str(row[correct_val_idx] or "").strip()
            notes = str(row[notes_idx] or "").strip()

            target_url = ""
            if correct_val:
                target_url = correct_val
            elif check_val == "Correct":
                target_url = orig_url

            if target_url:
                gold_rows.append({
                    "company_id": comp_id,
                    "category": category,
                    "correct_url": target_url,
                    "notes": notes,
                    "verified": "TRUE",
                })

        os.makedirs(os.path.dirname(out_path), exist_ok=True)
        with open(out_path, "w", newline="", encoding="utf-8-sig") as f:
            fieldnames = ["company_id", "category", "correct_url", "notes", "verified"]
            writer = csv.DictWriter(f, fieldnames=fieldnames)
            writer.writeheader()
            for r in gold_rows:
                writer.writerow(r)

        self.stdout.write(self.style.SUCCESS(f"Successfully extracted {len(gold_rows)} gold rows to: {out_path}"))

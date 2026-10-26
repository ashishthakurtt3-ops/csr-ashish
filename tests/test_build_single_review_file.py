import os
import csv
import shutil
import tempfile
import unittest
from unittest.mock import patch

import django
os.environ.setdefault("DJANGO_SETTINGS_MODULE", "config.settings")
django.setup()

from django.core.management import call_command
from django.core.management.base import CommandError
from django.db import connection

import openpyxl


class BuildSingleReviewFileTestCase(unittest.TestCase):
    """
    Unit test verifying build_single_review_file with a small fixture run folder:
    - Row ordering: COMPANY, PAGE, JOB, CSR, CONTACT, WARNING per company
    - NOT FOUND rows added for missing required page categories
    - VERIFY_THIS sample (all PAGE, all CSR, all CONTACT, first 5 JOBs, and any warned JOBs)
    - Clickable hyperlinks on url and website_url cells
    - Verification that the database is not touched (read-only command)
    """

    def setUp(self):
        # Create a self-contained temporary fixture run directory
        self.temp_dir = tempfile.mkdtemp(prefix="test_crawl_review_")

        # 1. Companies.csv (2 test companies)
        companies_file = os.path.join(self.temp_dir, "Companies.csv")
        with open(companies_file, "w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=[
                "company_id", "company_name", "website_url", "crawl_result", "reason",
                "time_taken_seconds", "pages_found", "jobs_found", "csr_items_found",
                "contacts_found", "ai_calls", "cache_hits"
            ])
            writer.writeheader()
            writer.writerow({
                "company_id": "TEST_COMP_01",
                "company_name": "Acme Innovations Ltd",
                "website_url": "https://www.acme-innovations.com",
                "crawl_result": "OK",
                "reason": "Both Jobs and CSR intelligence successfully extracted.",
                "time_taken_seconds": "15.2",
                "pages_found": "2",
                "jobs_found": "7",
                "csr_items_found": "2",
                "contacts_found": "1",
                "ai_calls": "2",
                "cache_hits": "1"
            })
            writer.writerow({
                "company_id": "TEST_COMP_02",
                "company_name": "Beta Global Technologies",
                "website_url": "https://www.betaglobal.com",
                "crawl_result": "PARTIAL",
                "reason": "Partial extraction: CSR portal not identified.",
                "time_taken_seconds": "22.0",
                "pages_found": "1",
                "jobs_found": "2",
                "csr_items_found": "0",
                "contacts_found": "1",
                "ai_calls": "0",
                "cache_hits": "0"
            })

        # 2. Pages.csv (Acme has JOBS and CSR; Beta has only JOBS)
        pages_file = os.path.join(self.temp_dir, "Pages.csv")
        with open(pages_file, "w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=[
                "company_id", "company_name", "category", "chosen_url", "alternatives",
                "how_found", "link_text", "menu_path", "ats_platform", "ats_tenant",
                "ownership_status", "validation_result", "ai_confidence", "ai_reason",
                "CHECK", "CORRECT_VALUE", "NOTES"
            ])
            writer.writeheader()
            writer.writerow({
                "company_id": "TEST_COMP_01",
                "company_name": "Acme Innovations Ltd",
                "category": "JOBS",
                "chosen_url": "https://careers.acme-innovations.com",
                "how_found": "HEADER_NAV",
                "link_text": "Work With Us",
                "menu_path": "Company > Careers",
                "ai_confidence": "0.95",
                "ai_reason": "Direct corporate ATS portal"
            })
            writer.writerow({
                "company_id": "TEST_COMP_01",
                "company_name": "Acme Innovations Ltd",
                "category": "CSR",
                "chosen_url": "https://www.acme-innovations.com/community-csr",
                "how_found": "FOOTER_NAV",
                "link_text": "CSR Initiatives",
                "menu_path": "About > CSR",
                "ai_confidence": "0.85",
                "ai_reason": "Statutory CSR disclosure"
            })
            writer.writerow({
                "company_id": "TEST_COMP_02",
                "company_name": "Beta Global Technologies",
                "category": "JOBS",
                "chosen_url": "https://www.betaglobal.com/careers",
                "how_found": "HEADER_NAV",
                "link_text": "Join Our Team",
                "menu_path": "About > Join",
                "ai_confidence": "0.70",
                "ai_reason": "Rules only fallback"
            })

        # 3. Jobs.csv (Acme has 7 jobs; job #7 has generic fallback warning; Beta has 2 jobs)
        jobs_file = os.path.join(self.temp_dir, "Jobs.csv")
        with open(jobs_file, "w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=[
                "company_id", "company_name", "job_title", "department", "location",
                "work_mode", "employment_type", "experience", "external_job_id",
                "ats_platform", "apply_source_url", "extraction_strategy_used",
                "description_snippet", "retrieved_at", "CHECK", "CORRECT_VALUE", "NOTES"
            ])
            writer.writeheader()
            # 6 regular jobs for Acme
            for idx in range(1, 7):
                writer.writerow({
                    "company_id": "TEST_COMP_01",
                    "company_name": "Acme Innovations Ltd",
                    "job_title": f"Software Engineer Tier {idx}",
                    "department": "Engineering",
                    "location": "Bengaluru",
                    "external_job_id": f"REQ_00{idx}",
                    "apply_source_url": f"https://careers.acme-innovations.com/jobs/REQ_00{idx}",
                    "extraction_strategy_used": "ats_api",
                    "description_snippet": f"Job snippet {idx}"
                })
            # 7th job for Acme has a generic listing link
            writer.writerow({
                "company_id": "TEST_COMP_01",
                "company_name": "Acme Innovations Ltd",
                "job_title": "Explore All Engineering Careers",
                "department": "Engineering",
                "location": "Remote",
                "external_job_id": "",
                "apply_source_url": "https://careers.acme-innovations.com/all-openings",
                "extraction_strategy_used": "html_table",
                "description_snippet": "Generic listing link"
            })
            # 2 jobs for Beta
            for idx in range(1, 3):
                writer.writerow({
                    "company_id": "TEST_COMP_02",
                    "company_name": "Beta Global Technologies",
                    "job_title": f"DevOps Specialist {idx}",
                    "department": "Cloud Infra",
                    "location": "Mumbai",
                    "external_job_id": f"BETA_{idx}",
                    "apply_source_url": f"https://www.betaglobal.com/careers/{idx}",
                    "extraction_strategy_used": "schema_org",
                    "description_snippet": f"DevOps description {idx}"
                })

        # 4. CSR.csv (Acme has 2 CSR projects; Beta has none)
        csr_file = os.path.join(self.temp_dir, "CSR.csv")
        with open(csr_file, "w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=[
                "company_id", "company_name", "project_or_disclosure_name", "theme",
                "financial_year", "budget", "geography", "source_url", "document_type",
                "page_number", "evidence_snippet", "CHECK", "CORRECT_VALUE", "NOTES"
            ])
            writer.writeheader()
            writer.writerow({
                "company_id": "TEST_COMP_01",
                "company_name": "Acme Innovations Ltd",
                "project_or_disclosure_name": "Clean Drinking Water Initiative",
                "theme": "Water & Sanitation",
                "financial_year": "2024-25",
                "budget": "5000000",
                "geography": "Karnataka",
                "source_url": "https://www.acme-innovations.com/community-csr/water",
                "document_type": "PDF",
                "page_number": "14",
                "evidence_snippet": "Project sanctioned for 50 lakh INR."
            })
            writer.writerow({
                "company_id": "TEST_COMP_01",
                "company_name": "Acme Innovations Ltd",
                "project_or_disclosure_name": "Primary Education Support",
                "theme": "Education",
                "financial_year": "2024-25",
                "budget": "3500000",
                "geography": "Maharashtra",
                "source_url": "https://www.acme-innovations.com/community-csr/education",
                "document_type": "web page",
                "page_number": "",
                "evidence_snippet": "School infrastructure and educational kits."
            })

        # 5. Contacts.csv (Acme has 1 contact; Beta has 1 contact)
        contacts_file = os.path.join(self.temp_dir, "Contacts.csv")
        with open(contacts_file, "w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=[
                "company_id", "company_name", "contact_type", "value", "department",
                "source_url", "evidence_snippet", "CHECK", "CORRECT_VALUE", "NOTES"
            ])
            writer.writeheader()
            writer.writerow({
                "company_id": "TEST_COMP_01",
                "company_name": "Acme Innovations Ltd",
                "contact_type": "email",
                "value": "csr@acme-innovations.com",
                "department": "CSR & Foundation",
                "source_url": "https://www.acme-innovations.com/contact",
                "evidence_snippet": "mailto:csr@acme-innovations.com"
            })
            writer.writerow({
                "company_id": "TEST_COMP_02",
                "company_name": "Beta Global Technologies",
                "contact_type": "phone",
                "value": "+91 22 12345678",
                "department": "Corporate Desk",
                "source_url": "https://www.betaglobal.com/reach-us",
                "evidence_snippet": "tel:+91 22 12345678"
            })

        # 6. Warnings.csv (Acme job #7 has GENERIC_LISTING_LINK warning; Beta has 1 standalone warning)
        warnings_file = os.path.join(self.temp_dir, "Warnings.csv")
        with open(warnings_file, "w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=[
                "company_id", "company_name", "warning_type", "details", "url", "evidence_snippet"
            ])
            writer.writeheader()
            writer.writerow({
                "company_id": "TEST_COMP_01",
                "company_name": "Acme Innovations Ltd",
                "warning_type": "GENERIC_LISTING_LINK",
                "details": "Job apply URL falls back to generic careers listing page.",
                "url": "https://careers.acme-innovations.com/all-openings",
                "evidence_snippet": "Explore All Engineering Careers"
            })
            writer.writerow({
                "company_id": "TEST_COMP_02",
                "company_name": "Beta Global Technologies",
                "warning_type": "UNRESOLVED_ATS_TENANT",
                "details": "Subdomain was found but multi-tenant organization ID could not be resolved.",
                "url": "https://betaglobal.taleo.net",
                "evidence_snippet": "Taleo portal requires client code"
            })

        # 7. README.csv
        readme_file = os.path.join(self.temp_dir, "README.csv")
        with open(readme_file, "w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)
            writer.writerow(["ImpactBridge Crawl Review Sampler", ""])
            writer.writerow(["Crawl Date & Time", "2026-10-23 11:27:55 UTC"])
            writer.writerow(["Page Classifier", "TwoStagePageClassifier"])

    def tearDown(self):
        shutil.rmtree(self.temp_dir, ignore_errors=True)

    def test_build_single_review_file_structure_and_verifications(self):
        """
        Verify that build_single_review_file produces company_review_single.xlsx with:
        - Exact row ordering: COMPANY, PAGE, JOB, CSR, CONTACT, WARNING per company
        - NOT FOUND rows for missing required page categories
        - Priority VERIFY_THIS sample (all PAGE, all CSR, all CONTACT, first 5 JOBs, and any warned JOBs)
        - Hyperlinks on valid URLs
        - Zero database queries executed
        """
        out_excel = os.path.join(self.temp_dir, "company_review_single.xlsx")
        self.assertFalse(os.path.exists(out_excel))

        # Check DB queries: read-only command must execute 0 DB queries
        initial_queries = len(connection.queries)
        with patch("django.db.backends.utils.CursorWrapper.execute") as mock_exec:
            call_command("build_single_review_file", run=self.temp_dir)
            mock_exec.assert_not_called()

        self.assertTrue(os.path.exists(out_excel), "Expected company_review_single.xlsx to be generated.")

        wb = openpyxl.load_workbook(out_excel)
        self.assertIn("Review", wb.sheetnames)
        self.assertIn("README", wb.sheetnames)

        ws_review = wb["Review"]
        headers = [c.value for c in ws_review[1]]

        expected_columns = [
            "company_no", "company_id", "company_name", "website_url", "record_type",
            "category", "title_or_name", "url", "how_found", "ats_platform_or_doc_type",
            "location_or_geography", "department_or_theme", "experience_or_FY", "id_or_budget",
            "ai_confidence_and_reason", "evidence_snippet", "warning", "VERIFY_THIS",
            "CHECK", "CORRECT_VALUE", "NOTES"
        ]
        self.assertEqual(headers, expected_columns)

        # Collect rows into list of dicts
        rows = []
        for r in ws_review.iter_rows(min_row=2, values_only=False):
            r_dict = {headers[idx]: r[idx] for idx in range(len(headers))}
            rows.append(r_dict)

        # ----------------------------------------------------------------------
        # 1. Verify Company 1 (TEST_COMP_01) Row Ordering & Content
        # ----------------------------------------------------------------------
        comp1_rows = [r for r in rows if r["company_id"].value == "TEST_COMP_01"]
        comp1_record_types = [r["record_type"].value for r in comp1_rows]

        # Order must be: 1 COMPANY, 5 PAGE (2 found + 3 NOT FOUND), 7 JOB, 2 CSR, 1 CONTACT, 0 WARNING (since generic warning is on JOB)
        expected_comp1_types = (
            ["COMPANY"] +
            ["PAGE"] * 5 +
            ["JOB"] * 7 +
            ["CSR"] * 2 +
            ["CONTACT"] * 1
        )
        self.assertEqual(comp1_record_types, expected_comp1_types)

        # Check NOT FOUND rows for Acme (missing SUSTAINABILITY, ANNUAL_REPORT, NGO_PARTNERSHIP)
        comp1_page_rows = [r for r in comp1_rows if r["record_type"].value == "PAGE"]
        comp1_not_found = [r for r in comp1_page_rows if r["title_or_name"].value == "NOT FOUND"]
        self.assertEqual(len(comp1_not_found), 3)
        not_found_categories = {r["category"].value for r in comp1_not_found}
        self.assertEqual(not_found_categories, {"SUSTAINABILITY", "ANNUAL_REPORT", "NGO_PARTNERSHIP"})
        for nf in comp1_not_found:
            self.assertEqual(nf["url"].value, None)
            self.assertEqual(nf["VERIFY_THIS"].value, "YES")

        # ----------------------------------------------------------------------
        # 2. Verify VERIFY_THIS Priority Sample on JOB Rows
        # (First 5 jobs + any job with a warning must have VERIFY_THIS = YES)
        # ----------------------------------------------------------------------
        comp1_job_rows = [r for r in comp1_rows if r["record_type"].value == "JOB"]
        self.assertEqual(len(comp1_job_rows), 7)

        # Jobs 1 to 5: must be YES
        for i in range(5):
            self.assertEqual(comp1_job_rows[i]["VERIFY_THIS"].value, "YES", f"Job index {i+1} should have VERIFY_THIS=YES")
            self.assertEqual(comp1_job_rows[i]["warning"].value, None)

        # Job 6 (index 5): regular job, beyond first 5, no warning -> must be blank
        self.assertEqual(comp1_job_rows[5]["VERIFY_THIS"].value, None)
        self.assertEqual(comp1_job_rows[5]["warning"].value, None)

        # Job 7 (index 6): generic listing warning -> must be marked warning and VERIFY_THIS=YES
        self.assertEqual(comp1_job_rows[6]["warning"].value, "GENERIC_LISTING_LINK")
        self.assertEqual(comp1_job_rows[6]["VERIFY_THIS"].value, "YES")

        # ----------------------------------------------------------------------
        # 3. Verify Company 2 (TEST_COMP_02) Row Ordering & Warning Row
        # ----------------------------------------------------------------------
        comp2_rows = [r for r in rows if r["company_id"].value == "TEST_COMP_02"]
        comp2_record_types = [r["record_type"].value for r in comp2_rows]

        # Order must be: 1 COMPANY, 5 PAGE (1 found + 4 NOT FOUND), 2 JOB, 0 CSR, 1 CONTACT, 1 WARNING
        expected_comp2_types = (
            ["COMPANY"] +
            ["PAGE"] * 5 +
            ["JOB"] * 2 +
            ["CONTACT"] * 1 +
            ["WARNING"] * 1
        )
        self.assertEqual(comp2_record_types, expected_comp2_types)

        # Standalone warning row check
        comp2_warn_row = [r for r in comp2_rows if r["record_type"].value == "WARNING"][0]
        self.assertEqual(comp2_warn_row["title_or_name"].value, "UNRESOLVED_ATS_TENANT")
        self.assertEqual(comp2_warn_row["warning"].value, "UNRESOLVED_ATS_TENANT")
        self.assertEqual(comp2_warn_row["VERIFY_THIS"].value, "YES")

        # ----------------------------------------------------------------------
        # 4. Verify Clickable Hyperlinks on url Cells
        # ----------------------------------------------------------------------
        url_cells_with_hyperlink = 0
        for r in rows:
            url_cell = r["url"]
            val_str = str(url_cell.value or "")
            if val_str.startswith("http://") or val_str.startswith("https://"):
                self.assertIsNotNone(url_cell.hyperlink, f"Expected clickable hyperlink on cell with URL {val_str}")
                self.assertEqual(url_cell.hyperlink.target, val_str)
                url_cells_with_hyperlink += 1
        self.assertGreater(url_cells_with_hyperlink, 10)

        # ----------------------------------------------------------------------
        # 5. Verify README sheet has title and company counts
        # ----------------------------------------------------------------------
        ws_readme = wb["README"]
        self.assertIn("ImpactBridge", str(ws_readme["A1"].value))
        # Look for Acme and Beta in the counts table
        readme_text = [str(cell.value) for row in ws_readme.iter_rows() for cell in row if cell.value]
        self.assertIn("Acme Innovations Ltd", readme_text)
        self.assertIn("Beta Global Technologies", readme_text)
